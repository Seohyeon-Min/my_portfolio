from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path('docs/Resume_Activision_Tech_Art.docx')
BLUE = '395BFF'
INK = '101827'
MUTED = '53627E'
LIGHT = 'F3F6FC'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.42)
sec.bottom_margin = Inches(0.40)
sec.left_margin = Inches(0.58)
sec.right_margin = Inches(0.58)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Aptos'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
normal.font.size = Pt(8.8)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.02

for name, size, bold, color in [('Resume Name', 23, True, INK), ('Resume Title', 10.5, True, BLUE), ('Resume Heading', 10.5, True, BLUE), ('Resume Role', 10.8, True, INK), ('Resume Meta', 8.3, False, MUTED)]:
    s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.font.name = 'Aptos Display' if name == 'Resume Name' else 'Aptos'
    s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name)
    s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name)
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.line_spacing = 1.0

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def borders(cell, color='D8DFED'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), color)
        tcBorders.append(el)

def cell_margin(cell, top=55, start=90, bottom=55, end=90):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_cell_text(cell, text, bold=False, color=INK, size=8.6):
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)

def section_heading(text):
    p=doc.add_paragraph(style='Resume Heading')
    p.paragraph_format.space_before=Pt(4.8); p.paragraph_format.space_after=Pt(2.2)
    p.add_run(text)
    pPr=p._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'10'); bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),BLUE); pbdr.append(bottom); pPr.append(pbdr)

def bullet(text):
    p=doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent=Inches(0.10); p.paragraph_format.first_line_indent=Inches(-0.10)
    p.paragraph_format.space_after=Pt(1.0)
    p.add_run('• ' + text)

def experience(name, date, role, details):
    t=doc.add_table(rows=1, cols=2); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width=Inches(5.15); t.columns[1].width=Inches(1.25)
    left,right=t.rows[0].cells
    left.width=Inches(5.15); right.width=Inches(1.25)
    for c in (left,right):
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=left.paragraphs[0]; p.style='Resume Role'; p.paragraph_format.space_before=Pt(2.1); p.paragraph_format.space_after=Pt(0); p.add_run(name)
    p=right.paragraphs[0]; p.style='Resume Title'; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before=Pt(2.1); p.add_run(date)
    p=doc.add_paragraph(style='Resume Meta'); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(1.4); p.add_run(role)
    for item in details: bullet(item)

# Header
table=doc.add_table(rows=1, cols=2); table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.columns[0].width=Inches(4.55); table.columns[1].width=Inches(2.35)
l,r=table.rows[0].cells; l.width=Inches(4.55); r.width=Inches(2.35)
p=l.paragraphs[0]; p.style='Resume Name'; p.add_run('MIN SEOHYEON')
p=l.add_paragraph(style='Resume Title'); p.add_run('TECHNICAL ARTIST | TOOLS, SHADERS & REAL-TIME CONTENT PIPELINES')
p=r.paragraphs[0]; p.style='Resume Meta'; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('weare1842@gmail.com\n'); p.add_run('github.com/Seohyeon-Min\n'); p.add_run('seohyeon-min.github.io/my_portfolio\n'); p.add_run('linkedin.com/in/seohyeon-min')

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(2.8); p.paragraph_format.line_spacing=1.05
p.add_run('Technical Artist ').bold=True
p.add_run('who builds artist-facing procedural tools, shaders, and real-time visual systems. Brings a strong 3D art foundation and engineering fluency to improve content-creation workflows, translate visual direction into reusable controls, and collaborate across art and programming.')

section_heading('TECHNICAL TOOLKIT')
t=doc.add_table(rows=4, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
t.columns[0].width=Inches(1.37); t.columns[1].width=Inches(5.53)
skills=[
    ('3D & Tech Art','Blender (modeling, sculpting, rigging, Geometry Nodes), procedural modeling, mesh Boolean/volume-remesh workflows, character integration, VFX'),
    ('Engines & Shaders','Unity URP, Unreal Engine, HLSL, ShaderLab, OpenGL, GLSL, Niagara/Cascade basics, RenderDoc, framebuffer post-processing'),
    ('Tools & Code','Python, C#, C++, object-oriented programming, Unreal editor tooling, DataAsset/ScriptableObject workflows, procedural animation'),
    ('Production','Perforce, Jira, Git, code review, debugging, profiling, Visual Studio, CMake; critique, technical feedback, visual integration')]
for row,(label,value) in zip(t.rows,skills):
    for c in row.cells:
        borders(c); cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(row.cells[0], LIGHT); set_cell_text(row.cells[0],label,True,INK,8.4); set_cell_text(row.cells[1],value,False,INK,8.45)

section_heading('TECHNICAL ART EXPERIENCE')
experience('CARBOOM','2026 - Present','Gameplay Programmer / Technical Art - Tools | Unreal Engine / Python / C++ / Perforce / Jira',[
    'Built an artist-facing Unreal Engine Python placement tool that distributes space-background planets from artist-editable DataAsset parameters, enabling rapid visual iteration without code changes.',
    'Translated composition intent - apparent size, density, and clustering - into reusable placement controls, bridging environment-art direction and procedural content generation.',
    'Exposed presets and editable controls that reduce engineering dependency during scene-composition iteration.'
])
experience('POSEIDON SKATE','Sep 2026','Technical Art / Shader Development / Character Modeling & Rigging | Unity URP / HLSL / C# / Blender / Jira / Perforce',[
    'Built ocean, wave, and tornado HLSL shaders using flow noise, domain warping, and Voronoi caustics for a rideable procedural water surface.',
    'Modeled and rigged a low-poly Poseidon character in Blender with a 25-bone skeleton, then integrated the FBX asset into Unity.',
    'Authored a landing-impact splash-ring shader with randomized per-bump timing and Voronoi facet detailing.'
])
experience('RUIN FORGE','2026 - Present','Technical Artist - Procedural Tools (Personal) | Blender Geometry Nodes / Python',[
    'Building a Geometry Nodes tool that lets environment artists reposition a single control object to regenerate layered concrete breaks, exposed inner surfaces, and crack placement instead of hand-modeling each variation.'
])
experience('STREET TYPER','Aug 2026','Technical Art / UI / Art | Unity URP / C# / ShaderLab / Notion',[
    'Owned original 2D art, UI composition, particles, outlines, camera shake, hit VFX, and animated feedback for a shipped bilingual typing-combat game.',
    'Specified, evaluated, debugged, and integrated a reusable UI shader workflow for rounded forms, gradients, shadows, blur, presets, and Inspector iteration.'
])
experience('TOO HOT!','Jul 2026','Technical Art / Visual Integration | Unity / ShaderLab / VFX / Notion',[
    'Created and integrated 2D shadow treatment, pattern-specific VFX, UI, animation, hit feedback, and visual hierarchy; tuned controls for readable shadows across combat spaces.',
    'Coordinated a 130+ item P0-P3 backlog, two programmers, code review, merges, and final visual integration.'
])
experience('NEW MANZO','Aug 2025 - Sep 2026','C# Programmer / Technical Art | Unity / C# / Notion',[
    'Implemented procedural leg animation for a multi-legged boss using ground raycasts and step-arc motion; built raycasting-based underwater visibility and post-processing.'
])
experience('MANZO','Sep 2024 - Dec 2025','Graphics / Engine Programmer | C++ / OpenGL / GLSL / Notion',[
    'Built a custom C++/OpenGL renderer with layer-sorted draw queues and framebuffer post-processing for bloom, underwater distortion, god rays, ripples, and transitions.',
    'Profiled severe frame drops, traced redundant per-frame collision checks, and removed repeated work to stabilize performance; largest repository contributor with 366 commits.'
])

section_heading('ADDITIONAL EXPERIENCE & EDUCATION')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1.5)
r=p.add_run('Teaching Assistant - Game Development Project I | Spring 2025  '); r.bold=True
p.add_run('Supported approximately 30 DigiPen Korea students with C++ implementation and debugging, providing actionable technical feedback on team projects.')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
r=p.add_run('B.S. Computer Science in Real-Time Interactive Simulation | Expected May 2028  '); r.bold=True
p.add_run('Keimyung University / DigiPen Institute of Technology | GPA 3.958/4.0')

doc.core_properties.title='Min Seohyeon Technical Artist Resume'
doc.core_properties.author='Min Seohyeon'
doc.save(OUT)
print(OUT)
